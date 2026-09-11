from __future__ import annotations

from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT / "drug_repurposing_deck"
OUTPUT_DECK = OUTPUT_DIR / "Drug_Repurposing_Agent_Prototype_Overview.pptx"
REFERENCE_DECK = Path(
    "/Users/zhongqishuai/Documents/materials/applied_ai_decks/deck_ready_for_sharing/"
    "Cloudera AI Deck for Agentic AI Bigger Picture -- Manufacturing Example (ZH).pptx"
)

NAVY = "171242"
NAVY_2 = "241B57"
PURPLE = "5B50F6"
PURPLE_LIGHT = "EEECFF"
BLUE_LIGHT = "EEF4FF"
ORANGE = "FF5A1F"
ORANGE_LIGHT = "FFF1EB"
GREEN = "00A878"
GREEN_LIGHT = "E8F7F2"
AMBER = "E88B20"
AMBER_LIGHT = "FFF4E5"
INK = "202136"
MUTED = "667085"
LIGHT = "F6F7FB"
MID = "D8DCE8"
WHITE = "FFFFFF"

FONT = "Plus Jakarta Sans"
FONT_LIGHT = "Plus Jakarta Sans Light"
FONT_MEDIUM = "Plus Jakarta Sans Medium"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_exact_typeface(run, font: str) -> None:
    """Set every OOXML script slot, matching the reference deck's font wiring."""
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs", "a:sym"):
        existing = rpr.find(qn(tag))
        if existing is not None:
            rpr.remove(existing)
        node = OxmlElement(tag)
        node.set("typeface", font)
        rpr.append(node)


def apply_box_shadow(shape) -> None:
    """Reference standard: black shadow, 10% opacity, 135°, 3pt distance, 6pt blur."""
    sppr = shape._element.spPr
    old = sppr.find(qn("a:effectLst"))
    if old is not None:
        sppr.remove(old)
    effects = OxmlElement("a:effectLst")
    shadow = OxmlElement("a:outerShdw")
    shadow.set("blurRad", str(6 * 12700))
    shadow.set("dist", str(3 * 12700))
    shadow.set("dir", str(135 * 60000))
    shadow.set("rotWithShape", "0")
    color = OxmlElement("a:srgbClr")
    color.set("val", "000000")
    alpha = OxmlElement("a:alpha")
    alpha.set("val", "10000")
    color.append(alpha)
    shadow.append(color)
    effects.append(shadow)
    sppr.append(effects)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 10,
    color: str = INK,
    font: str = FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0,
    line_spacing: float = 1.05,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for index, line in enumerate(text.split("\n")):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.line_spacing = line_spacing
        paragraph.space_before = paragraph.space_after = Pt(0)
        for run in paragraph.runs:
            set_exact_typeface(run, font)
            run.font.size = Pt(size)
            run.font.bold = False
            run.font.color.rgb = rgb(color)
    return box


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str = WHITE,
    line: str | None = MID,
    line_width: float = 0.8,
    radius: bool = True,
):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    apply_box_shadow(shape)
    return shape


def add_line(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str = PURPLE,
    width: float = 1.4,
    arrow: bool = True,
    dashed: bool = False,
):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if dashed:
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        line_xml = shape._element.spPr.ln
        end = line_xml.makeelement(qn("a:tailEnd"))
        end.set("type", "triangle")
        line_xml.append(end)
    return shape


def add_badge(slide, text: str, x: float, y: float, w: float, *, fill: str, color: str):
    add_rect(slide, x, y, w, 0.27, fill=fill, line=None)
    add_text(
        slide,
        text,
        x,
        y + 0.008,
        w,
        0.23,
        size=7.2,
        color=color,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_header(slide, title: str, subtitle: str, motif_blob: bytes):
    add_text(slide, title, 0.58, 0.25, 8.32, 0.48, size=23, color=NAVY, font=FONT)
    add_text(slide, subtitle, 0.58, 0.76, 8.32, 0.42, size=11, color=PURPLE, font=FONT_MEDIUM)
    slide.shapes.add_picture(BytesIO(motif_blob), Inches(9.08), Inches(0), width=Inches(0.92))


def add_footer(slide, logo_blob: bytes):
    add_rect(slide, 0, 5.54, 10, 0.08, fill=ORANGE, line=None, radius=False)
    add_text(slide, "©2026 Cloudera, Inc. All Rights Reserved.", 0.42, 5.20, 4.6, 0.16, size=5.1, color="A6A6B2", font=FONT_LIGHT)
    slide.shapes.add_picture(BytesIO(logo_blob), Inches(8.27), Inches(5.20), width=Inches(1.22))


def add_number(slide, value: str, x: float, y: float, *, fill: str):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.28), Inches(0.28))
    circle.fill.solid()
    circle.fill.fore_color.rgb = rgb(fill)
    circle.line.fill.background()
    apply_box_shadow(circle)
    add_text(slide, value, x, y + 0.004, 0.28, 0.24, size=7.7, color=WHITE, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def build_overview(slide, motif_blob: bytes, logo_blob: bytes):
    add_header(
        slide,
        "Drug Repurposing Discovery | Prototype at a Glance",
        "A compact agentic prototype that connects a repurposing question to explicit graph evidence and an AI-generated scientific report.",
        motif_blob,
    )

    add_rect(slide, 0.58, 1.28, 8.82, 0.60, fill=NAVY, line=None)
    add_text(slide, "SHOWCASE QUESTION", 0.78, 1.42, 1.24, 0.16, size=7.2, color="BEB8FF")
    add_text(slide, "Could Semaglutide treat Obesity?", 2.14, 1.35, 4.65, 0.29, size=16.5, color=WHITE)
    add_badge(slide, "CLINICAL VALIDATION REQUIRED", 7.32, 1.43, 1.70, fill="3A316E", color=WHITE)

    cards = [
        (
            0.58,
            "01",
            "EVIDENCE",
            "Seeded knowledge graph",
            "24 entities · 32 directed relationships\nFive knowledge-source domains\nOptional publication triplet preview",
            PURPLE,
            PURPLE_LIGHT,
        ),
        (
            3.55,
            "02",
            "GRAPH ANALYSIS",
            "Deterministic path selection",
            "Fixed showcase entity pair\nBFS evaluates four candidate paths\nConfidence, hidden links + formula score",
            ORANGE,
            ORANGE_LIGHT,
        ),
        (
            6.52,
            "03",
            "AGENT OUTPUT",
            "Grounded discovery report",
            "Highlighted mechanism path\nHypothesis, safety rationale + significance\nRisks, next steps + SSE activity",
            GREEN,
            GREEN_LIGHT,
        ),
    ]
    for x, number, label, heading, body, accent, pale in cards:
        add_rect(slide, x, 2.08, 2.72, 2.34, fill=WHITE, line=MID)
        add_rect(slide, x, 2.08, 2.72, 0.07, fill=accent, line=None, radius=False)
        add_number(slide, number, x + 0.18, 2.30, fill=accent)
        add_text(slide, label, x + 0.55, 2.365, 1.72, 0.15, size=7.2, color=accent)
        add_text(slide, heading, x + 0.18, 2.78, 2.34, 0.48, size=12.0, color=NAVY)
        add_text(slide, body, x + 0.18, 3.34, 2.34, 0.82, size=7.8, color=INK, font=FONT_LIGHT, line_spacing=1.20)

    add_line(slide, 3.30, 3.22, 3.50, 3.22, color=ORANGE, width=2.0)
    add_line(slide, 6.27, 3.22, 6.47, 3.22, color=GREEN, width=2.0)
    add_rect(slide, 0.58, 4.62, 8.82, 0.38, fill=LIGHT, line=None)
    add_text(
        slide,
        "Graph tools select the path and compute the numeric scores; Claude interprets that evidence for people.",
        0.78,
        4.72,
        8.42,
        0.18,
        size=8.5,
        color=NAVY,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, logo_blob)


def build_architecture(slide, motif_blob: bytes, logo_blob: bytes):
    add_header(
        slide,
        "Architecture | Graph Evidence Grounds the Agent",
        "The implemented path is small and inspectable: JSON evidence, deterministic graph analytics, one Claude agent, and a React discovery experience.",
        motif_blob,
    )

    columns = [0.42, 2.30, 5.10, 7.42]
    widths = [1.48, 2.35, 1.86, 2.12]
    labels = ["INPUTS", "GRAPH ANALYTICS", "DISCOVERY AGENT", "EXPERIENCE"]
    colors = [PURPLE, ORANGE, PURPLE, GREEN]
    fills = [PURPLE_LIGHT, ORANGE_LIGHT, PURPLE_LIGHT, GREEN_LIGHT]
    for x, w, label, color, fill in zip(columns, widths, labels, colors, fills):
        add_badge(slide, label, x, 1.24, w, fill=fill, color=color)

    add_rect(slide, 0.42, 1.72, 1.48, 2.36, fill=LIGHT, line=MID)
    add_text(slide, "Question", 0.62, 1.96, 1.08, 0.20, size=10.5, color=NAVY, align=PP_ALIGN.CENTER)
    add_text(slide, "Semaglutide\n→ Obesity", 0.62, 2.34, 1.08, 0.54, size=9.2, color=PURPLE, align=PP_ALIGN.CENTER, line_spacing=1.12)
    add_rect(slide, 0.61, 3.17, 1.10, 0.55, fill=WHITE, line=PURPLE)
    add_text(slide, "seed_graph.json\n24 nodes · 32 edges", 0.69, 3.27, 0.94, 0.34, size=6.7, color=INK, font=FONT_LIGHT, align=PP_ALIGN.CENTER)

    add_rect(slide, 2.30, 1.72, 2.35, 2.36, fill=ORANGE_LIGHT, line=ORANGE)
    steps = [
        ("1", "Directed BFS", "Find simple paths, max depth 10"),
        ("2", "Rank candidates", "Average confidence, then path length"),
        ("3", "Compute metrics", "Hidden-link count + opportunity score"),
    ]
    for index, (number, title, detail) in enumerate(steps):
        y = 1.96 + index * 0.62
        add_number(slide, number, 2.52, y, fill=ORANGE)
        add_text(slide, title, 2.91, y + 0.005, 1.45, 0.18, size=8.7, color=NAVY)
        add_text(slide, detail, 2.91, y + 0.23, 1.45, 0.25, size=6.5, color=MUTED, font=FONT_LIGHT)
    add_badge(slide, "TOP PATH → STRUCTURED JSON", 2.61, 3.66, 1.73, fill=WHITE, color=ORANGE)

    add_rect(slide, 5.10, 1.72, 1.86, 2.36, fill=WHITE, line=PURPLE)
    add_text(slide, "Claude Sonnet", 5.30, 1.98, 1.46, 0.24, size=10.8, color=NAVY, align=PP_ALIGN.CENTER)
    add_text(slide, "question + path JSON", 5.30, 2.40, 1.46, 0.18, size=7.2, color=PURPLE, align=PP_ALIGN.CENTER)
    add_text(slide, "Hypothesis\nMechanism explanation\nSafety rationale\nKnowledge fragmentation\nRisks + next steps", 5.30, 2.78, 1.46, 0.94, size=7.0, color=INK, font=FONT_LIGHT, align=PP_ALIGN.CENTER, line_spacing=1.14)

    add_rect(slide, 7.42, 1.72, 2.12, 2.36, fill=GREEN_LIGHT, line=GREEN)
    add_text(slide, "Flask + SSE + React", 7.62, 1.98, 1.72, 0.24, size=10.3, color=NAVY, align=PP_ALIGN.CENTER)
    add_text(slide, "Live activity feed", 7.62, 2.43, 1.72, 0.18, size=7.4, color=GREEN, align=PP_ALIGN.CENTER)
    add_text(slide, "Force-directed graph\nHighlighted evidence path\nStructured discovery report\nClinical validation warning", 7.62, 2.82, 1.72, 0.82, size=7.2, color=INK, font=FONT_LIGHT, align=PP_ALIGN.CENTER, line_spacing=1.15)

    add_line(slide, 1.94, 2.90, 2.25, 2.90, color=PURPLE, width=2.0)
    add_line(slide, 4.69, 2.90, 5.05, 2.90, color=PURPLE, width=2.0)
    add_line(slide, 7.00, 2.90, 7.37, 2.90, color=GREEN, width=2.0)

    add_rect(slide, 0.58, 4.42, 8.82, 0.52, fill=LIGHT, line=MID)
    add_text(slide, "OPTIONAL PUBLICATION PATH", 0.78, 4.58, 1.46, 0.16, size=6.8, color=MUTED)
    add_text(slide, "TXT / decoded text", 2.40, 4.56, 1.10, 0.18, size=7.3, color=NAVY, align=PP_ALIGN.CENTER)
    add_line(slide, 3.54, 4.66, 4.02, 4.66, color=MUTED, width=1.0, dashed=True)
    add_text(slide, "Claude triplet extraction", 4.09, 4.56, 1.48, 0.18, size=7.3, color=NAVY, align=PP_ALIGN.CENTER)
    add_line(slide, 5.62, 4.66, 6.10, 4.66, color=MUTED, width=1.0, dashed=True)
    add_text(slide, "Preview in UI · no graph write-back", 6.18, 4.56, 2.51, 0.18, size=7.3, color=ORANGE, align=PP_ALIGN.CENTER)
    add_footer(slide, logo_blob)


def add_path_node(slide, label: str, x: float, y: float, *, fill: str, line: str, color: str = NAVY):
    add_rect(slide, x, y, 1.60, 0.48, fill=fill, line=line)
    add_text(slide, label, x + 0.08, y + 0.09, 1.44, 0.24, size=7.6, color=color, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def build_path(slide, motif_blob: bytes, logo_blob: bytes):
    add_header(
        slide,
        "Discovery Path | From Semaglutide to Obesity",
        "The selected route is the highest-confidence of four candidates; one cross-domain connection is explicitly flagged as hidden knowledge.",
        motif_blob,
    )

    top = [
        ("Semaglutide", PURPLE_LIGHT, PURPLE),
        ("GLP-1 Receptor", BLUE_LIGHT, PURPLE),
        ("Hypothalamus", LIGHT, MID),
        ("Arcuate Nucleus", LIGHT, MID),
        ("NPY Neurons", AMBER_LIGHT, AMBER),
    ]
    top_x = [0.45, 2.22, 3.99, 5.76, 7.53]
    for x, (label, fill, line) in zip(top_x, top):
        add_path_node(slide, label, x, 1.44, fill=fill, line=line)
    top_relations = ["agonizes", "expressed_in", "contains", "contains"]
    for idx, relation in enumerate(top_relations):
        color = ORANGE if relation == "expressed_in" else PURPLE
        add_line(slide, top_x[idx] + 1.60, 1.68, top_x[idx + 1] - 0.04, 1.68, color=color, width=1.6, dashed=(relation == "expressed_in"))
    add_badge(slide, "HIDDEN", 3.12, 1.91, 0.64, fill=ORANGE_LIGHT, color=ORANGE)

    bottom = [
        ("Food Intake", GREEN_LIGHT, GREEN),
        ("Energy Homeostasis", ORANGE_LIGHT, ORANGE),
        ("Body Weight", GREEN_LIGHT, GREEN),
        ("Obesity", NAVY, NAVY),
    ]
    bottom_x = [7.53, 5.50, 3.47, 1.44]
    add_line(slide, 8.33, 1.95, 8.33, 2.30, color=PURPLE, width=1.6)
    for x, (label, fill, line) in zip(bottom_x, bottom):
        add_path_node(slide, label, x, 2.34, fill=fill, line=line, color=WHITE if fill == NAVY else NAVY)
    bottom_relations = ["affects", "determines", "diagnostic_for"]
    for idx, relation in enumerate(bottom_relations):
        add_line(slide, bottom_x[idx] - 0.04, 2.58, bottom_x[idx + 1] + 1.60, 2.58, color=PURPLE, width=1.6)
    add_text(
        slide,
        "agonizes  →  expressed_in [hidden]  →  contains  →  contains  →  enhances  →  affects  →  determines  →  diagnostic_for",
        0.58,
        2.92,
        8.82,
        0.16,
        size=6.0,
        color=PURPLE,
        font=FONT_LIGHT,
        align=PP_ALIGN.CENTER,
    )

    stats = [
        ("4", "candidate paths"),
        ("8", "edges selected"),
        ("93.6%", "avg. edge confidence"),
        ("1", "hidden connection"),
        ("0.77", "opportunity score"),
    ]
    for idx, (value, label) in enumerate(stats):
        x = 0.58 + idx * 1.77
        add_rect(slide, x, 3.18, 1.59, 0.74, fill=WHITE, line=MID)
        add_text(slide, value, x + 0.08, 3.30, 1.43, 0.23, size=12.2, color=ORANGE if idx == 4 else NAVY, align=PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.08, 3.59, 1.43, 0.15, size=6.2, color=MUTED, font=FONT_LIGHT, align=PP_ALIGN.CENTER)

    add_rect(slide, 0.58, 4.17, 4.22, 0.69, fill=ORANGE_LIGHT, line=ORANGE)
    add_text(slide, "DETERMINISTIC RESULT", 0.78, 4.31, 1.38, 0.15, size=6.8, color=ORANGE)
    add_text(slide, "Top path + confidence + hidden-link count + score", 0.78, 4.56, 3.80, 0.18, size=7.6, color=NAVY)
    add_rect(slide, 5.18, 4.17, 4.22, 0.69, fill=PURPLE_LIGHT, line=PURPLE)
    add_text(slide, "CLAUDE SYNTHESIS", 5.38, 4.31, 1.30, 0.15, size=6.8, color=PURPLE)
    add_text(slide, "Hypothesis + mechanism + safety + risks + next steps", 5.38, 4.56, 3.80, 0.18, size=7.6, color=NAVY)
    add_text(
        slide,
        "Current boundary: fixed entity pair · seeded in-memory JSON · Claude API required · no publication write-back",
        0.58,
        4.98,
        8.82,
        0.17,
        size=6.6,
        color=MUTED,
        font=FONT_LIGHT,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, logo_blob)


def extract_brand_assets() -> tuple[bytes, bytes]:
    reference = Presentation(REFERENCE_DECK)
    source_slide = reference.slides[15]
    pictures = [shape for shape in source_slide.shapes if shape.shape_type == 13]
    motif = min(pictures, key=lambda shape: shape.top)
    logo = max(pictures, key=lambda shape: shape.top)
    return motif.image.blob, logo.image.blob


def validate(prs: Presentation) -> None:
    assert len(prs.slides) == 3
    assert prs.slide_width == Inches(10)
    assert prs.slide_height == Inches(5.625)
    for slide_number, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            assert shape.left >= 0, (slide_number, shape.name, "left")
            assert shape.top >= 0, (slide_number, shape.name, "top")
            assert shape.left + shape.width <= prs.slide_width + 1, (slide_number, shape.name, "right")
            assert shape.top + shape.height <= prs.slide_height + 1, (slide_number, shape.name, "bottom")
        text_shapes = [shape for shape in slide.shapes if getattr(shape, "has_text_frame", False)]
        assert text_shapes[0].text_frame.paragraphs[0].runs[0].font.name == FONT
        assert text_shapes[0].text_frame.paragraphs[0].runs[0].font.size == Pt(23)
        assert text_shapes[1].text_frame.paragraphs[0].runs[0].font.name == FONT_MEDIUM
        assert text_shapes[1].text_frame.paragraphs[0].runs[0].font.size == Pt(11)
        for shape in text_shapes[2:]:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    assert run.font.name in {FONT, FONT_LIGHT}, (slide_number, shape.text, run.font.name)
                    assert run.font.bold is False
                    rpr = run._r.get_or_add_rPr()
                    for tag in ("a:latin", "a:ea", "a:cs", "a:sym"):
                        node = rpr.find(qn(tag))
                        assert node is not None and node.get("typeface") == run.font.name, (slide_number, shape.text, tag)
        for shape in slide.shapes:
            if shape.shape_type == 1:
                effects = shape._element.spPr.find(qn("a:effectLst"))
                shadow = effects.find(qn("a:outerShdw")) if effects is not None else None
                assert shadow is not None, (slide_number, shape.name, "shadow")
                assert shadow.get("blurRad") == str(6 * 12700)
                assert shadow.get("dist") == str(3 * 12700)
                assert shadow.get("dir") == str(135 * 60000)
                color = shadow.find(qn("a:srgbClr"))
                alpha = color.find(qn("a:alpha")) if color is not None else None
                assert alpha is not None and alpha.get("val") == "10000"


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    motif_blob, logo_blob = extract_brand_assets()
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    blank = prs.slide_layouts[6]
    builders = [build_overview, build_architecture, build_path]
    for builder in builders:
        slide = prs.slides.add_slide(blank)
        builder(slide, motif_blob, logo_blob)
    validate(prs)
    prs.save(OUTPUT_DECK)
    print(OUTPUT_DECK)


if __name__ == "__main__":
    main()
