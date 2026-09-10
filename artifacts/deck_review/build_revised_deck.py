from __future__ import annotations

from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
SOURCE_DECK = ROOT / "input" / "main.pptx"
SOURCE_SCREENSHOT = ROOT / "input" / "demo-dashboard.png"
OUTPUT_DECK = ROOT / "Cloudera_AI_Deck_Manufacturing_ZH_revised.pptx"
ASSET_DIR = ROOT / "generated_assets"

NAVY = "171242"
NAVY_2 = "241B57"
PURPLE = "5B50F6"
PURPLE_LIGHT = "EEECFF"
ORANGE = "FF5A1F"
ORANGE_LIGHT = "FFF1EB"
GREEN = "00A878"
GREEN_LIGHT = "E8F7F2"
AMBER = "E88B20"
AMBER_LIGHT = "FFF4E5"
INK = "202136"
MUTED = "667085"
LIGHT = "F6F7FB"
MID = "D8DCE8"
WHITE = "FFFFFF"

FONT = "Plus Jakarta Sans"
FONT_MEDIUM = "Plus Jakarta Sans Medium"
FONT_SEMIBOLD = "Plus Jakarta Sans SemiBold"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def clear_slide(slide) -> None:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 11,
    color: str = INK,
    font: str = FONT,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.0,
    line_spacing: float = 1.05,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    tf.word_wrap = True
    lines = text.split("\n")
    for index, line in enumerate(lines):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        for run in p.runs:
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
    return box


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str = WHITE,
    line: str | None = MID,
    line_width: float = 0.8,
    radius: bool = True,
):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    return shape


def add_line(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str = PURPLE,
    width: float = 1.5,
    arrow: bool = True,
    dashed: bool = False,
):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    if dashed:
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        ln = line._element.spPr.ln
        tail_end = ln.makeelement(qn("a:tailEnd"))
        tail_end.set("type", "triangle")
        ln.append(tail_end)
    return line


def add_badge(slide, text: str, x: float, y: float, w: float, *, fill: str, color: str):
    add_rect(slide, x, y, w, 0.28, fill=fill, line=None)
    add_text(
        slide,
        text,
        x,
        y + 0.01,
        w,
        0.24,
        size=7.5,
        color=color,
        font=FONT_SEMIBOLD,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )


def add_header(slide, title: str, subtitle: str, motif_blob: bytes):
    add_text(slide, title, 0.58, 0.28, 8.15, 0.43, size=19, color=NAVY, font=FONT_MEDIUM)
    add_text(slide, subtitle, 0.58, 0.73, 8.20, 0.42, size=8.7, color=PURPLE)
    slide.shapes.add_picture(BytesIO(motif_blob), Inches(9.08), Inches(0), width=Inches(0.92))


def add_footer(slide, logo_blob: bytes):
    add_rect(slide, 0, 5.54, 10, 0.08, fill=ORANGE, line=None, radius=False)
    add_text(slide, "©2026 Cloudera, Inc. All Rights Reserved.", 0.42, 5.20, 4.6, 0.16, size=5.1, color="A6A6B2")
    slide.shapes.add_picture(BytesIO(logo_blob), Inches(8.27), Inches(5.20), width=Inches(1.22))


def add_number(slide, value: str, x: float, y: float, *, fill: str = PURPLE):
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.28), Inches(0.28))
    circle.fill.solid()
    circle.fill.fore_color.rgb = rgb(fill)
    circle.line.fill.background()
    add_text(slide, value, x, y + 0.005, 0.28, 0.25, size=8, color=WHITE, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)


def build_slide_15(slide, motif_blob: bytes, logo_blob: bytes):
    add_header(
        slide,
        "演示总览｜新设计的电池能否通过阶段评审？",
        "一个构建并运行于 Cloudera AI Workbench 的车辆 NPI 原型，把工程数据、ML 风险与历史知识汇成可审计的人工评审材料。",
        motif_blob,
    )

    add_rect(slide, 0.58, 1.22, 8.82, 0.62, fill=NAVY, line=None)
    add_text(slide, "业务问题", 0.78, 1.36, 0.82, 0.2, size=8, color="BEB8FF", font=FONT_SEMIBOLD, bold=True)
    add_text(slide, "PACK-ATLAS-01 是否可以放行？", 1.62, 1.28, 5.1, 0.34, size=17, color=WHITE, font=FONT_SEMIBOLD, bold=True)
    add_badge(slide, "人工最终决策", 7.63, 1.38, 1.42, fill="3A316E", color=WHITE)

    cards = [
        (
            0.58,
            "01  输入",
            "项目事实",
            "• 设计需求与参数\n• BOM 与供应商数据\n• 当前与历史 DVP&R 结果\n• 历史设计图像与报告",
            PURPLE,
            PURPLE_LIGHT,
        ),
        (
            3.55,
            "02  分析",
            "ML + Agentic AI",
            "• 确定性测试台与追溯矩阵\n• GBM 设计风险评分\n• 5 个专业 Agent 并行取证\n• Lance 多模态知识检索",
            ORANGE,
            ORANGE_LIGHT,
        ),
        (
            6.52,
            "03  输出",
            "可审计评审材料",
            "• PASS / CONDITIONAL / FAIL\n• 风险、缺口与标准覆盖\n• 引用证据与叙述性总结\n• 人工批准 / 附条件 / 驳回",
            GREEN,
            GREEN_LIGHT,
        ),
    ]
    for x, label, heading, body, accent, pale in cards:
        add_rect(slide, x, 2.04, 2.72, 2.30, fill=WHITE, line=MID)
        add_rect(slide, x, 2.04, 2.72, 0.07, fill=accent, line=None, radius=False)
        add_badge(slide, label, x + 0.18, 2.25, 0.82, fill=pale, color=accent)
        add_text(slide, heading, x + 0.18, 2.65, 2.32, 0.36, size=12.5, color=NAVY, font=FONT_SEMIBOLD, bold=True)
        add_text(slide, body, x + 0.18, 3.08, 2.32, 1.05, size=8.2, color=INK, line_spacing=1.18)

    add_line(slide, 3.30, 3.18, 3.50, 3.18, color=ORANGE, width=2.2)
    add_line(slide, 6.27, 3.18, 6.47, 3.18, color=GREEN, width=2.2)
    add_rect(slide, 0.58, 4.54, 8.82, 0.45, fill=LIGHT, line=None)
    add_text(slide, "适合参会者", 0.78, 4.66, 0.92, 0.19, size=7.6, color=MUTED, font=FONT_SEMIBOLD, bold=True)
    add_text(slide, "业务负责人  ·  工程与项目团队  ·  数据与 AI 开发团队", 1.75, 4.60, 4.65, 0.27, size=10.2, color=NAVY, font=FONT_MEDIUM)
    add_badge(slide, "代码决定，LLM 叙述", 7.28, 4.62, 1.77, fill=ORANGE_LIGHT, color=ORANGE)
    add_footer(slide, logo_blob)


def build_slide_16(slide, motif_blob: bytes, logo_blob: bytes):
    add_header(
        slide,
        "端到端架构｜一套数据底座，两条分析路径，一个人工决策",
        "工程应用与 Agent 调用同一套受治理的数据与 API；所有分数、测试结论和建议都可重复计算。",
        motif_blob,
    )

    columns = [0.45, 2.45, 5.08, 7.55]
    widths = [1.62, 2.18, 2.02, 1.98]
    heads = ["工程数据", "确定性分析 + ML", "Agentic AI 协调", "评审输出"]
    colors = [PURPLE, ORANGE, PURPLE, GREEN]
    for x, w, head, color in zip(columns, widths, heads, colors):
        add_badge(slide, head, x, 1.22, w, fill=(PURPLE_LIGHT if color == PURPLE else ORANGE_LIGHT if color == ORANGE else GREEN_LIGHT), color=color)

    add_rect(slide, 0.45, 1.68, 1.62, 2.80, fill=LIGHT, line=MID)
    add_text(slide, "PACK-ATLAS-01", 0.63, 1.88, 1.26, 0.28, size=10, color=NAVY, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "需求\n设计参数\nBOM / 供应商\n测试计划与结果", 0.66, 2.36, 1.20, 1.42, size=8.6, color=INK, align=PP_ALIGN.CENTER, line_spacing=1.25)
    add_badge(slide, "CSV → Iceberg / SDX", 0.64, 4.04, 1.24, fill=WHITE, color=MUTED)

    add_rect(slide, 2.45, 1.68, 2.18, 1.05, fill=ORANGE_LIGHT, line=ORANGE)
    add_text(slide, "测试台 + 追溯矩阵", 2.64, 1.87, 1.80, 0.24, size=10.5, color=NAVY, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "覆盖率 80%  ·  热测试 MARGINAL", 2.64, 2.23, 1.80, 0.20, size=7.2, color=ORANGE, font=FONT_MEDIUM, align=PP_ALIGN.CENTER)
    add_rect(slide, 2.45, 2.91, 2.18, 1.05, fill=ORANGE_LIGHT, line=ORANGE)
    add_text(slide, "GBM 设计风险模型", 2.64, 3.10, 1.80, 0.24, size=10.5, color=NAVY, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "热需求 HIGH  ·  特征贡献可解释", 2.64, 3.46, 1.80, 0.20, size=7.2, color=ORANGE, font=FONT_MEDIUM, align=PP_ALIGN.CENTER)
    add_rect(slide, 2.45, 4.12, 2.18, 0.36, fill=PURPLE_LIGHT, line=PURPLE)
    add_text(slide, "Lance：文本 + 图像 + 报告 + 元数据", 2.57, 4.21, 1.94, 0.16, size=6.7, color=PURPLE, font=FONT_MEDIUM, align=PP_ALIGN.CENTER)

    add_rect(slide, 5.08, 1.68, 2.02, 2.80, fill=WHITE, line=PURPLE)
    add_text(slide, "5 个专业 Agent", 5.26, 1.86, 1.66, 0.28, size=11, color=NAVY, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    for idx, label in enumerate(["Coverage", "Test Verdict", "Compliance", "Design Risk", "Knowledge Reuse"]):
        y = 2.30 + idx * 0.37
        add_rect(slide, 5.34, y, 1.50, 0.27, fill=LIGHT, line=None)
        add_text(slide, label, 5.45, y + 0.045, 1.28, 0.15, size=6.9, color=INK, font=FONT_MEDIUM, align=PP_ALIGN.CENTER)
    add_text(slide, "查询 · 取证 · 汇总", 5.26, 4.22, 1.66, 0.16, size=7.2, color=PURPLE, font=FONT_MEDIUM, align=PP_ALIGN.CENTER)

    add_rect(slide, 7.55, 1.68, 1.98, 0.94, fill=NAVY, line=None)
    add_text(slide, "traceability.decide()", 7.71, 1.88, 1.66, 0.19, size=8.4, color="C7C2FF", font=FONT_MEDIUM, align=PP_ALIGN.CENTER)
    add_text(slide, "CONDITIONAL", 7.71, 2.14, 1.66, 0.25, size=12, color=WHITE, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, 7.55, 2.86, 1.98, 0.70, fill=PURPLE_LIGHT, line=PURPLE)
    add_text(slide, "LLM 生成评审叙述\n仅影响文字", 7.74, 3.01, 1.60, 0.35, size=7.8, color=NAVY, font=FONT_MEDIUM, align=PP_ALIGN.CENTER)
    add_rect(slide, 7.55, 3.80, 1.98, 0.68, fill=GREEN_LIGHT, line=GREEN)
    add_text(slide, "人工最终裁定\n批准 / 附条件 / 驳回", 7.72, 3.93, 1.64, 0.36, size=7.8, color=NAVY, font=FONT_MEDIUM, align=PP_ALIGN.CENTER)

    add_line(slide, 2.08, 3.08, 2.40, 3.08, color=PURPLE, width=2.0)
    add_line(slide, 4.66, 3.08, 5.03, 3.08, color=PURPLE, width=2.0)
    add_line(slide, 7.12, 3.08, 7.50, 3.08, color=PURPLE, width=2.0)
    add_rect(slide, 0.58, 4.72, 8.80, 0.30, fill=ORANGE_LIGHT, line=None)
    add_text(slide, "两条独立信号相互印证：ML 预测 HIGH  +  实测 MARGINAL  →  代码给出 CONDITIONAL。", 0.78, 4.79, 8.40, 0.16, size=8.0, color=ORANGE, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, logo_blob)


def make_dashboard_crops() -> list[Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    image = Image.open(SOURCE_SCREENSHOT).convert("RGB")
    crops = [
        (0, 92, 596, 865),
        (582, 92, 1332, 1060),
        (1330, 92, 1915, 1060),
    ]
    paths: list[Path] = []
    for index, box in enumerate(crops, 1):
        path = ASSET_DIR / f"dashboard_crop_{index}.png"
        image.crop(box).save(path, optimize=True)
        paths.append(path)
    return paths


def add_picture_contain(slide, path: Path, x: float, y: float, w: float, h: float):
    with Image.open(path) as im:
        ratio = im.width / im.height
    box_ratio = w / h
    if ratio >= box_ratio:
        draw_w = w
        draw_h = w / ratio
        draw_x = x
        draw_y = y + (h - draw_h) / 2
    else:
        draw_h = h
        draw_w = h * ratio
        draw_x = x + (w - draw_w) / 2
        draw_y = y
    return slide.shapes.add_picture(str(path), Inches(draw_x), Inches(draw_y), width=Inches(draw_w), height=Inches(draw_h))


def build_slide_17(slide, motif_blob: bytes, logo_blob: bytes, crops: list[Path]):
    add_header(
        slide,
        "项目驾驶舱｜风险、证据与结论在一屏对齐",
        "演示从左到右展开：先看 ML 风险，再核对独立测试证据，最后复用历史多模态知识。",
        motif_blob,
    )
    panels = [
        (0.44, "01", "ML 预警", "热需求被评为 HIGH", ORANGE, ORANGE_LIGHT),
        (3.55, "02", "工程实证", "热测试 MARGINAL · 覆盖率 80%", AMBER, AMBER_LIGHT),
        (6.66, "03", "知识复用", "Lance 检索历史图像与测试报告", PURPLE, PURPLE_LIGHT),
    ]
    for index, (x, num, heading, caption, accent, pale) in enumerate(panels):
        add_rect(slide, x, 1.24, 2.90, 3.72, fill=WHITE, line=MID)
        add_rect(slide, x, 1.24, 2.90, 0.06, fill=accent, line=None, radius=False)
        add_number(slide, num, x + 0.16, 1.43, fill=accent)
        add_text(slide, heading, x + 0.52, 1.43, 2.08, 0.25, size=10.5, color=NAVY, font=FONT_SEMIBOLD, bold=True)
        add_rect(slide, x + 0.14, 1.82, 2.62, 2.36, fill=LIGHT, line="E4E6EF")
        add_picture_contain(slide, crops[index], x + 0.19, 1.87, 2.52, 2.26)
        add_rect(slide, x + 0.14, 4.34, 2.62, 0.42, fill=pale, line=None)
        add_text(slide, caption, x + 0.22, 4.43, 2.46, 0.19, size=7.5, color=accent, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "讲解顺序：① 左侧 ML 给出前瞻风险  →  ② 中间测试与追溯给出独立事实  →  ③ 右侧知识库补充历史上下文  →  点击 Run Gate Review", 0.58, 5.08, 8.82, 0.22, size=7.2, color=MUTED, font=FONT_MEDIUM, align=PP_ALIGN.CENTER)
    add_footer(slide, logo_blob)


def build_slide_18(slide, motif_blob: bytes, logo_blob: bytes):
    add_header(
        slide,
        "Gate Review｜五个专业 Agent 并行取证，代码统一裁决",
        "Agent 负责查询、解释和引用证据；决定函数只读取确定性结果，LLM 不参与评分或控制流。",
        motif_blob,
    )
    add_rect(slide, 0.44, 2.18, 1.46, 1.00, fill=NAVY, line=None)
    add_text(slide, "Run Gate Review", 0.59, 2.44, 1.16, 0.21, size=10.5, color=WHITE, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "SSE 实时推送", 0.59, 2.76, 1.16, 0.17, size=6.7, color="C7C2FF", align=PP_ALIGN.CENTER)

    agents = [
        ("Coverage", "覆盖率与缺口"),
        ("Test Verdict", "DVP&R 结论"),
        ("Compliance", "标准覆盖"),
        ("Design Risk", "ML 风险解释"),
        ("Knowledge Reuse", "历史多模态证据"),
    ]
    for idx, (name, desc) in enumerate(agents):
        y = 1.30 + idx * 0.72
        add_rect(slide, 2.40, y, 2.34, 0.52, fill=WHITE, line=PURPLE)
        add_text(slide, name, 2.56, y + 0.09, 1.12, 0.17, size=8.6, color=NAVY, font=FONT_SEMIBOLD, bold=True)
        add_text(slide, desc, 3.62, y + 0.10, 0.96, 0.16, size=6.8, color=MUTED, align=PP_ALIGN.RIGHT)
        add_line(slide, 1.90, 2.68, 2.35, y + 0.26, color=PURPLE, width=1.1, arrow=True)

    add_rect(slide, 5.23, 1.75, 1.34, 2.23, fill=LIGHT, line=MID)
    add_text(slide, "证据包", 5.43, 1.96, 0.94, 0.22, size=10.5, color=NAVY, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "结构化发现\n来源 ID\n参数与阈值\n证据哈希", 5.43, 2.42, 0.94, 1.05, size=8.1, color=INK, align=PP_ALIGN.CENTER, line_spacing=1.20)
    for idx in range(5):
        y = 1.30 + idx * 0.72 + 0.26
        add_line(slide, 4.77, y, 5.18, 2.86, color=PURPLE, width=1.0, arrow=True)

    add_rect(slide, 7.03, 1.61, 2.34, 1.16, fill=NAVY, line=None)
    add_text(slide, "traceability.decide()", 7.25, 1.83, 1.90, 0.19, size=8.4, color="C7C2FF", font=FONT_MEDIUM, align=PP_ALIGN.CENTER)
    add_text(slide, "CONDITIONAL", 7.25, 2.14, 1.90, 0.28, size=14, color=WHITE, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    add_line(slide, 6.60, 2.86, 6.98, 2.19, color=ORANGE, width=2.0)

    add_rect(slide, 7.03, 3.08, 1.08, 1.02, fill=PURPLE_LIGHT, line=PURPLE)
    add_text(slide, "LLM", 7.20, 3.29, 0.74, 0.20, size=10, color=NAVY, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "写叙述", 7.20, 3.61, 0.74, 0.16, size=7.2, color=PURPLE, align=PP_ALIGN.CENTER)
    add_rect(slide, 8.29, 3.08, 1.08, 1.02, fill=GREEN_LIGHT, line=GREEN)
    add_text(slide, "人工", 8.46, 3.29, 0.74, 0.20, size=10, color=NAVY, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "最终裁定", 8.46, 3.61, 0.74, 0.16, size=7.2, color=GREEN, align=PP_ALIGN.CENTER)
    add_line(slide, 8.20, 2.79, 7.57, 3.03, color=PURPLE, width=1.2)
    add_line(slide, 8.20, 2.79, 8.83, 3.03, color=GREEN, width=1.2)

    add_badge(slide, "相同输入 = 相同结果", 0.58, 4.77, 1.78, fill=ORANGE_LIGHT, color=ORANGE)
    add_badge(slide, "无 LLM 也能完整运行", 2.52, 4.77, 1.92, fill=PURPLE_LIGHT, color=PURPLE)
    add_text(slide, "所有建议都可回放到原始数据、阈值、模型版本与引用资产。", 4.68, 4.80, 4.54, 0.18, size=7.6, color=MUTED, font=FONT_MEDIUM, align=PP_ALIGN.RIGHT)
    add_footer(slide, logo_blob)


def build_slide_19(slide, motif_blob: bytes, logo_blob: bytes):
    add_header(
        slide,
        "一个原型，展示四类可复用能力",
        "从工程数据到人工关口决策：传统 ML、Agentic AI 与多模态知识索引在同一个 Cloudera AI 应用中协同。",
        motif_blob,
    )
    cards = [
        (0.45, "01", "数据与工程知识", "需求 · 设计参数 · BOM · 测试\n端到端追溯与缺口识别\n统一 API，面向真实 PLM / MES / ERP", PURPLE, PURPLE_LIGHT),
        (2.70, "02", "传统 ML + 生成式 AI", "GBM 为每条需求评分\nAgent 解释模型与工程事实\nLLM 可替换，且不改变任何结论", ORANGE, ORANGE_LIGHT),
        (4.95, "03", "多模态知识索引", "Lance 统一索引文本、图像与报告\n本体连接项目、需求、测试与资产\n历史经验可检索、可引用、可复用", GREEN, GREEN_LIGHT),
        (7.20, "04", "Agentic 协作与治理", "5 个专业 Agent 并行取证\nSSE 实时过程 + 哈希证据记录\n模板回退与人工最终决策", NAVY, "ECEBF5"),
    ]
    for x, num, title, body, accent, pale in cards:
        add_rect(slide, x, 1.33, 2.05, 2.72, fill=WHITE, line=MID)
        add_rect(slide, x, 1.33, 2.05, 0.07, fill=accent, line=None, radius=False)
        add_number(slide, num, x + 0.16, 1.59, fill=accent)
        add_text(slide, title, x + 0.16, 2.03, 1.72, 0.49, size=10.6, color=NAVY, font=FONT_SEMIBOLD, bold=True)
        add_text(slide, body, x + 0.16, 2.65, 1.72, 1.05, size=7.5, color=INK, line_spacing=1.20)
        add_rect(slide, x + 0.16, 3.71, 1.72, 0.18, fill=pale, line=None)

    add_rect(slide, 0.45, 4.27, 8.80, 0.47, fill=NAVY, line=None)
    add_text(slide, "Cloudera AI Workbench", 0.68, 4.39, 1.62, 0.19, size=8.0, color=WHITE, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "CAII / vLLM / Ollama", 2.52, 4.39, 1.56, 0.19, size=8.0, color="C7C2FF", font=FONT_MEDIUM, align=PP_ALIGN.CENTER)
    add_text(slide, "FastAPI + Swagger", 4.30, 4.39, 1.48, 0.19, size=8.0, color="C7C2FF", font=FONT_MEDIUM, align=PP_ALIGN.CENTER)
    add_text(slide, "Agent Studio / MCP ready", 6.00, 4.39, 1.72, 0.19, size=8.0, color="C7C2FF", font=FONT_MEDIUM, align=PP_ALIGN.CENTER)
    add_text(slide, "React 项目驾驶舱", 7.94, 4.39, 1.08, 0.19, size=8.0, color=WHITE, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "落地路径：把合成数据源替换为企业数据源与生产模型目录，保留 API、Agent 流程和前端体验。", 0.58, 4.94, 8.82, 0.22, size=7.6, color=ORANGE, font=FONT_SEMIBOLD, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, logo_blob)


def main() -> None:
    assert SOURCE_DECK.exists(), SOURCE_DECK
    assert SOURCE_SCREENSHOT.exists(), SOURCE_SCREENSHOT

    prs = Presentation(str(SOURCE_DECK))
    assert len(prs.slides) >= 19
    assert round(prs.slide_width / 914400, 2) == 10.00
    assert round(prs.slide_height / 914400, 2) == 5.62

    source_slide = prs.slides[15]
    motif_blob = source_slide.shapes[18].image.blob
    logo_blob = source_slide.shapes[21].image.blob
    crops = make_dashboard_crops()

    builders = [
        lambda slide: build_slide_15(slide, motif_blob, logo_blob),
        lambda slide: build_slide_16(slide, motif_blob, logo_blob),
        lambda slide: build_slide_17(slide, motif_blob, logo_blob, crops),
        lambda slide: build_slide_18(slide, motif_blob, logo_blob),
        lambda slide: build_slide_19(slide, motif_blob, logo_blob),
    ]
    for slide_number, builder in zip(range(15, 20), builders):
        slide = prs.slides[slide_number - 1]
        clear_slide(slide)
        builder(slide)

    prs.save(str(OUTPUT_DECK))
    check = Presentation(str(OUTPUT_DECK))
    assert len(check.slides) == len(prs.slides)
    for slide_number in range(15, 20):
        text = " ".join(shape.text for shape in check.slides[slide_number - 1].shapes if hasattr(shape, "text"))
        assert text.strip()
        assert "Microsoft YaHei" not in text
    assert "新设计的电池能否通过阶段评审" in " ".join(
        shape.text for shape in check.slides[14].shapes if hasattr(shape, "text")
    )
    print(OUTPUT_DECK)


if __name__ == "__main__":
    main()
