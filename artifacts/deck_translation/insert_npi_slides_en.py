from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re

from pptx import Presentation
from pptx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
SOURCE_ZH = ROOT / "input" / "source_zh.pptx"
TARGET_EN = ROOT / "input" / "target_en.pptx"
OUTPUT = ROOT / "Cloudera_AI_Deck_Manufacturing_EN_with_NPI.pptx"


# Each entry maps a source shape index to translated paragraph content. A paragraph
# may be a string or a list of run strings when the source uses multiple styles.
TRANSLATIONS: dict[int, dict[int, list[str | list[str]]]] = {
    15: {
        0: [["Autonomous Gate Review on Your Data", " (R&D Example)"]],
        1: ["Example: From a program at release gate to an auditable PASS / CONDITIONAL / FAIL decision—all on your data, with no egress."],
        4: [["GATE REVIEW", "FLOW"]],
        6: ["01 · INGEST", "Program enters gate"],
        8: ["The PACK-ATLAS-01 battery-pack program enters the release gate, triggering a five-agent review"],
        10: [["02 · ", "ENGINEERING AGENT"], "Coverage & test verdicts"],
        12: ["Builds the traceability matrix—one requirement lacks a test plan; thermal test is MARGINAL at +3% margin"],
        14: ["03 · RISK & COMPLIANCE", "Scores & standards"],
        16: ["The design-risk model flags thermal as HIGH; checks coverage for UN 38.3 · GB 38031 · IEC 62660"],
        18: ["04 · KNOWLEDGE REUSE", "Historical validation"],
        20: [
            "Retrieves PACK-ORION-00 thermal design image + test report (historical FAIL)",
            "That failure corroborates HIGH risk + MARGINAL testing—confirming the concern",
        ],
        22: ["RESULT", "Recommendation + evidence"],
        24: ["CONDITIONAL—computed by traceability.decide(); findings form an evidence chain for human review."],
        25: [[
            "A deterministic core plus a connected knowledge layer turns raw program data into",
            " ",
            "a trusted, auditable gate decision—on your data.",
        ]],
    },
    16: {
        0: [["Can the New Battery Design Pass the Gate Review", "?"]],
        1: ["A vehicle NPI prototype built and running in Cloudera AI Workbench, turning engineering data, ML risk and historical knowledge into auditable materials for human review."],
        4: ["BUSINESS QUESTION"],
        5: ["Can PACK-ATLAS-01 be released?"],
        7: ["01  INPUT"],
        8: ["Program facts"],
        9: [
            "• Design requirements & parameters",
            "• BOM & supplier data",
            "• Current & historical DVP&R results",
            "• Historical design images & reports",
        ],
        11: ["02  ANALYSIS"],
        12: ["ML + Agentic AI"],
        13: [
            "• Deterministic test bench & traceability",
            "• GBM design-risk scoring",
            "• 5 specialist agents gather evidence in parallel",
            "• Lance multimodal knowledge retrieval",
        ],
        15: ["03  OUTPUT"],
        16: ["Auditable review package"],
        17: [
            "• PASS / CONDITIONAL / FAIL",
            "• Risks, gaps & standards coverage",
            "• Cited evidence & narrative summary",
            "• Human approve / condition / reject",
        ],
        21: ["Code decides. The LLM narrates."],
    },
    17: {
        0: ["End-to-End Architecture | Two Paths, One Decision"],
        1: ["Engineering apps and agents call the same governed data and APIs; every score, test verdict and recommendation is reproducible."],
        3: ["ENGINEERING DATA"],
        4: ["DETERMINISTIC ANALYTICS + ML"],
        5: ["AGENTIC AI ORCHESTRATION"],
        6: ["REVIEW OUTPUT"],
        8: ["PACK-ATLAS-01"],
        9: ["Requirements", "Design parameters", "BOM / suppliers", "Test plans & results"],
        10: ["Iceberg / SDX"],
        12: ["Test bench + traceability matrix"],
        13: ["80% coverage  ·  Thermal test MARGINAL"],
        15: ["GBM design-risk model"],
        16: [["Predictive risk", " ·  explainable feature contributions"]],
        18: ["Lance: text + images + reports + metadata"],
        20: ["5 specialist agents"],
        22: ["Coverage"],
        24: ["Test Verdict"],
        26: ["Compliance"],
        28: ["Design Risk"],
        30: ["Knowledge Reuse"],
        31: ["Query · evidence · synthesis"],
        33: ["LLM writes the review narrative", "Text only"],
        35: ["Human final decision", "Approve / condition / reject"],
        40: ["Two independent signals corroborate: ML predicts HIGH + test measures MARGINAL → code returns CONDITIONAL."],
    },
    18: {
        0: ["Risk, Evidence and Decision—Aligned on One Screen"],
        1: ["Move left to right: start with ML risk, verify independent test evidence, then reuse historical multimodal knowledge."],
        4: ["01"],
        5: [["ML ", "PREDICTIVE ", "RISK"]],
        7: ["Thermal requirement scored HIGH"],
        9: ["02"],
        10: ["ENGINEERING EVIDENCE"],
        12: ["Thermal test MARGINAL · 80% coverage"],
        14: ["03"],
        15: ["KNOWLEDGE REUSE"],
        17: ["Lance retrieves historical images and test reports"],
        18: ["Talk track: ① ML surfaces predictive risk → ② tests and traceability provide independent facts → ③ the knowledge base adds historical context → click Run Gate Review"],
    },
    19: {
        0: ["Gate Review | Five Agents Gather Evidence; Code Decides"],
        1: ["Agents query, explain and cite evidence; the decision function reads deterministic results only. The LLM never scores or controls the flow."],
        4: ["Run Gate Review"],
        5: ["SSE live updates"],
        7: ["Coverage"],
        8: ["Coverage & gaps"],
        11: ["Test Verdict"],
        12: ["DVP&R verdicts"],
        15: ["Compliance"],
        16: ["Standards coverage"],
        19: ["Design Risk"],
        20: ["ML risk explanation"],
        23: ["Knowledge Reuse"],
        24: ["Historical multimodal evidence"],
        27: ["Evidence package"],
        28: ["Structured findings", "Source IDs", "Parameters & thresholds", "Evidence hashes"],
        35: ["traceability.decide()"],
        36: ["CONDITIONAL"],
        39: ["LLM"],
        40: ["Writes narrative"],
        42: ["HUMAN"],
        43: ["Final decision"],
        47: ["Same input = Same result"],
        48: ["Every recommendation can be replayed to its source data, thresholds, model version and cited assets."],
    },
    20: {
        0: ["One Prototype, Four Reusable Capabilities"],
        1: ["From engineering data to a human gate decision: traditional ML, Agentic AI and multimodal knowledge indexing work together in one Cloudera AI application."],
        5: ["01"],
        6: ["DATA & ENGINEERING KNOWLEDGE"],
        7: [
            "Requirements · design parameters · BOM · tests",
            "End-to-end traceability & gap detection",
            "One API for real PLM / MES / ERP",
        ],
        10: ["02"],
        11: ["TRADITIONAL ML + GEN AI"],
        12: [
            "GBM scores risk per requirement",
            "Agents explain model and engineering evidence",
            "The LLM is swappable and never changes decisions",
        ],
        15: ["03"],
        16: ["MULTIMODAL KNOWLEDGE INDEX"],
        17: [
            "Lance indexes text, images and reports",
            "Ontology links programs, requirements, tests and assets",
            "Past experience is searchable, citable and reusable",
        ],
        20: ["04"],
        21: ["AGENTIC COLLABORATION & GOVERNANCE"],
        22: [
            "5 specialist agents gather evidence in parallel",
            "SSE live flow + hashed evidence records",
            "Template fallback + human final decision",
        ],
        23: ["Production path: replace synthetic data with enterprise sources and the production model registry; retain the API, agent workflow and frontend experience."],
    },
}


REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
CJK_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff]")


def layout_for_source_slide(target_prs: Presentation, source_slide):
    source_layout_name = str(source_slide.slide_layout.part.partname).rsplit("/", 1)[-1]
    for layout in target_prs.slide_layouts:
        if str(layout.part.partname).rsplit("/", 1)[-1] == source_layout_name:
            return layout
    raise KeyError(f"No matching target layout for {source_layout_name}")


def clear_shapes(slide) -> None:
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def clone_shape(source_shape, source_slide, target_slide) -> None:
    element = deepcopy(source_shape.element)
    for node in element.iter():
        for attr_name, old_rid in list(node.attrib.items()):
            if not attr_name.startswith(REL_NS):
                continue
            rel = source_slide.part.rels[old_rid]
            if rel.reltype.endswith("/image"):
                _, new_rid = target_slide.part.get_or_add_image_part(BytesIO(rel.target_part.blob))
                node.set(attr_name, new_rid)
            elif rel.is_external:
                new_rid = target_slide.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
                node.set(attr_name, new_rid)
            else:
                raise ValueError(f"Unsupported shape relationship: {rel.reltype}")
    target_slide.shapes._spTree.insert_element_before(element, "p:extLst")


def set_paragraph_text(paragraph, translation: str | list[str]) -> None:
    if isinstance(translation, list):
        if len(translation) != len(paragraph.runs):
            raise ValueError(
                f"Run count mismatch for {paragraph.text!r}: expected {len(paragraph.runs)}, got {len(translation)}"
            )
        for run, text in zip(paragraph.runs, translation):
            run.text = text
        return

    if not paragraph.runs:
        paragraph.text = translation
        return
    paragraph.runs[0].text = translation
    for run in paragraph.runs[1:]:
        run.text = ""


def translate_slide(slide_number: int, slide) -> None:
    mapping = TRANSLATIONS[slide_number]
    visible_text_shape_indices = {
        index for index, shape in enumerate(slide.shapes)
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    }
    if visible_text_shape_indices != set(mapping):
        missing = visible_text_shape_indices - set(mapping)
        extra = set(mapping) - visible_text_shape_indices
        raise ValueError(f"Slide {slide_number} translation mismatch; missing={missing}, extra={extra}")

    for shape_index, paragraphs in mapping.items():
        shape = slide.shapes[shape_index]
        if len(shape.text_frame.paragraphs) != len(paragraphs):
            raise ValueError(
                f"Slide {slide_number}, shape {shape_index}: paragraph count mismatch "
                f"({len(shape.text_frame.paragraphs)} != {len(paragraphs)})"
            )
        for paragraph, translation in zip(shape.text_frame.paragraphs, paragraphs):
            set_paragraph_text(paragraph, translation)


def shape_signature(slide):
    return [
        (
            str(shape.shape_type),
            shape.left,
            shape.top,
            shape.width,
            shape.height,
            shape.rotation,
        )
        for shape in slide.shapes
    ]


def insert_slide_ids(prs: Presentation, inserted_slides: list, position: int) -> None:
    slide_id_list = prs.slides._sldIdLst
    new_ids = list(slide_id_list[-len(inserted_slides):])
    for slide_id in new_ids:
        slide_id_list.remove(slide_id)
    for offset, slide_id in enumerate(new_ids):
        slide_id_list.insert(position + offset, slide_id)


def main() -> None:
    assert SOURCE_ZH.exists(), SOURCE_ZH
    assert TARGET_EN.exists(), TARGET_EN
    source = Presentation(str(SOURCE_ZH))
    target = Presentation(str(TARGET_EN))
    assert len(source.slides) == 41
    assert len(target.slides) == 35
    assert source.slide_width == target.slide_width
    assert source.slide_height == target.slide_height

    original_target_slide_15_text = " ".join(
        shape.text for shape in target.slides[14].shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )
    assert "IT runs it" in original_target_slide_15_text

    inserted = []
    for source_slide_number in range(15, 21):
        source_slide = source.slides[source_slide_number - 1]
        target_slide = target.slides.add_slide(layout_for_source_slide(target, source_slide))
        clear_shapes(target_slide)
        for shape in source_slide.shapes:
            clone_shape(shape, source_slide, target_slide)
        assert shape_signature(source_slide) == shape_signature(target_slide)
        translate_slide(source_slide_number, target_slide)
        inserted.append(target_slide)

    insert_slide_ids(target, inserted, position=14)
    target.save(str(OUTPUT))

    check = Presentation(str(OUTPUT))
    assert len(check.slides) == 41
    for source_slide_number in range(15, 21):
        source_slide = source.slides[source_slide_number - 1]
        inserted_slide = check.slides[source_slide_number - 1]
        assert shape_signature(source_slide) == shape_signature(inserted_slide)
        visible_text = " ".join(
            shape.text for shape in inserted_slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        )
        assert not CJK_RE.search(visible_text), (source_slide_number, visible_text)

    shifted_slide_text = " ".join(
        shape.text for shape in check.slides[20].shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    )
    assert "IT runs it" in shifted_slide_text
    print(OUTPUT)


if __name__ == "__main__":
    main()
